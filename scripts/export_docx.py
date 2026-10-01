"""Export the thesis chapters (Markdown) to Word (.docx) with native Word equations.

    pip install pypandoc_binary python-docx
    python scripts/export_docx.py            # writes docs/word/*.docx

Equations become editable Word equations, figures are embedded, tables get
borders, and body text is set to Times New Roman 12 pt (adjust FONT below to
match the department's template, or paste the content into the template).
"""

import re
import sys
import tempfile
from pathlib import Path

import _bootstrap  # noqa: F401

DOCS = Path("docs")
OUT = DOCS / "word"
CHAPTERS = ["Chapter3_Theoretical_Considerations.md", "Chapter4_Design_Considerations.md", "Appendix_Rule_Tables.md"]
FONT, SIZE_PT = "Times New Roman", 12


def prepare(md: str) -> str:
    # Word equations have no \tag: render the number at the right instead
    return re.sub(r"\\tag\{([^}]*)\}", r"\\qquad\\qquad (\1)", md)


def _add_borders(table) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "808080")
        borders.append(el)
    tbl_pr.append(borders)


def style(path: Path) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document(str(path))
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in [s.name for s in doc.styles]:
            st = doc.styles[name]
            st.font.name = FONT
            st.font.size = Pt(SIZE_PT)
    for table in doc.tables:
        _add_borders(table)
        for row in table.rows:
            for cell in row.cells:
                for par in cell.paragraphs:
                    for run in par.runs:
                        run.font.size = Pt(10)
                        run.font.name = FONT
    doc.save(str(path))


def main() -> int:
    import pypandoc

    OUT.mkdir(parents=True, exist_ok=True)
    for name in CHAPTERS:
        src = DOCS / name
        with tempfile.NamedTemporaryFile("w", suffix=".md", delete=False, encoding="utf-8") as tmp:
            tmp.write(prepare(src.read_text(encoding="utf-8")))
        dest = OUT / src.with_suffix(".docx").name
        pypandoc.convert_file(tmp.name, "docx", format="markdown+tex_math_dollars+pipe_tables-implicit_figures",
                              outputfile=str(dest), extra_args=[f"--resource-path={DOCS}"])
        Path(tmp.name).unlink()
        try:
            style(dest)
        except ImportError:
            print("python-docx not installed: skipped font/table styling")
        print(f"wrote {dest}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
